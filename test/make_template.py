# يبني نسخة اختبار تحاكي بنية قالب قيود HR الرسمي (للاختبار فقط — القوائم الحقيقية تُقرأ من القالب الرسمي)
import openpyxl, datetime
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter as L

heads = ["first_name","last_name","first_name_ar","last_name_ar","gender","date_of_birth","marital_status","nationality","residence_country","national_id_number","national_id_issued_date","national_id_expiry_date","email","job_title","work_phone_number","employment_status","department","employee_no","manager","location","work_email","home_phone_number","address","emergency_contact_name","emergency_contact_relationship","emergency_contact_home_phone_number","emergency_contact_mobile_phone_number","join_date","employment_type","contract_type","contract_end_date","attendance_setting","education_level","education_description","bank_name","bank_code","account_number","bank_account_currency","bank_address","iban","swift_code","salary_start_date","salary_basic_salary_amount","salary_pay_frequency","salary_basic_salary_kind","salary_input_time_type"]
assert len(heads)==46
req = set(heads[:17])|{"join_date","employment_type"}
countries = "السعودية;الإمارات العربيّة المتّحدة;قطر;الكويت;البحرين;عمان;العراق;مصر;الأردن;لبنان;فلسطين;سوريا;اليمن;المغرب;تونس;الجزائر;ليبيا;السودان;الهند;باكستان;بنغلاديش;الفلبّين;سريلانكا;نيبال;إندونيسيا;إثيوبيا;كينيا;نيجيريا;تركيّا;أفغانستان;المملكة المتحدة;الولايات المتّحدة;كندا;فرنسا;ألمانيا;الاتّحاد الرّوسي;الصّين;اليابان;كوريا، جمهوريّة كوريا;إيران، الجمهوريّة الإسلاميّة الإيرانيّة".split(";")
currencies = "ريال سعودي;درهم إماراتي;ريال قطري;دينار كويتي;دينار بحريني;ريال عماني;دولار أمريكي;يورو;جنيه إسترليني;جنيه مصري;دينار أردني;روبية هندية;روبية باكستانية;بيزو فلبيني".split(";")
lists = {
 "gender":["ذكر","أنثى"],
 "marital_status":["أعزب","متزوج","مطلق","أرمل"],
 "nationality":countries,
 "residence_country":countries,
 "employment_status":["مفعّل","غير مفعّل","فترة الإشعار"],
 "department":["الموارد البشرية","المالية","الهندسة","المبيعات","التسويق","العمليات"],
 "location":["المكتب الرئيسي","فرع الرياض"],
 "employment_type":["دوام كامل","دوام جزئي","عقد"],
 "contract_type":["محدد المدة","غير محدد المدة"],
 "attendance_setting":["افتراضي"],
 "education_level":["أقل من الثانوية","شهادة ثانوية","درجة المنتسبين","درجة البكالوريوس","درجة ماجستير","درجة دكتوراه"],
 "bank_account_currency":currencies,
 "salary_pay_frequency":["Daily","Weekly","Monthly"],
 "salary_basic_salary_kind":["Fixed","Hourly"],
 "salary_input_time_type":["وقت التوظيف","وقت العمل الفعلي"],
}
wb = openpyxl.Workbook()
ws = wb.active; ws.title="Employees"
ref = wb.create_sheet("Reference")
for i,h in enumerate(heads,1): ws.cell(1,i,h+(" *" if h in req else ""))
sample = {"first_name":"Ahmed","last_name":"Hassan","first_name_ar":"أحمد","last_name_ar":"حسن","gender":"ذكر","date_of_birth":datetime.date(1990,1,1),"marital_status":"أعزب","nationality":"السعودية","residence_country":"السعودية","national_id_number":"1012345678","national_id_issued_date":datetime.date(2020,1,1),"national_id_expiry_date":datetime.date(2029,12,31),"email":"ahmed.hassan@example.com","job_title":"Software Engineer","work_phone_number":"+966500000000","employment_status":"مفعّل","department":"الموارد البشرية","location":"المكتب الرئيسي","address":"King Fahd Road, Riyadh","join_date":datetime.date(2024,1,1),"employment_type":"دوام كامل","contract_type":"غير محدد المدة","attendance_setting":"افتراضي","education_level":"درجة البكالوريوس","bank_account_currency":"ريال سعودي","salary_start_date":datetime.date(2024,1,1),"salary_basic_salary_amount":10000,"salary_pay_frequency":"Monthly","salary_basic_salary_kind":"Fixed","salary_input_time_type":"وقت التوظيف"}
for i,h in enumerate(heads,1):
    v = sample.get(h)
    if v is not None:
        c = ws.cell(2,i,v)
        if isinstance(v, datetime.date): c.number_format="yyyy-mm-dd"
refcols = list(lists.keys())
for j,k in enumerate(refcols,1):
    ref.cell(1,j,k)
    for i,v in enumerate(lists[k],2): ref.cell(i,j,v)
col = {h:i+1 for i,h in enumerate(heads)}
for j,k in enumerate(refcols,1):
    dv = DataValidation(type="list", formula1=f"Reference!${L(j)}$2:${L(j)}${len(lists[k])+1}", allow_blank=True)
    ws.add_data_validation(dv); dv.add(f"{L(col[k])}2:{L(col[k])}1000")
# علامة "قالب تجريبي": الأداة تحذّر إذا استُخدم للتصدير الحقيقي (قوائمه ليست قوائم أي شركة)
wb.properties.title = "QHR_SAMPLE_TEMPLATE"
wb.properties.subject = "قالب تجريبي للاختبار فقط - لا ترفعه على قيود"
wb.save("template_test.xlsx")
print("ok", len(heads), len(req))

# نسخة بقوائم أغنى (أقسام وفروع واقعية) لملفات الاختبار الموسّعة ولتجربة المستخدم
import sys
if "--rich" in sys.argv:
    rich = dict(lists)
    rich["department"] = ["الموارد البشرية","المالية","الهندسة","المبيعات","التسويق","العمليات","تقنية المعلومات","خدمة العملاء","المشتريات","الشؤون الإدارية"]
    rich["location"] = ["المقر الرئيسي - الرياض","فرع جدة","فرع الدمام"]
    rich["attendance_setting"] = ["افتراضي","دوام مسائي"]
    for j,k in enumerate(refcols,1):
        for r in range(2, ref.max_row+1): ref.cell(r,j).value = None
        for i,v in enumerate(rich[k],2): ref.cell(i,j,v)
    ws.data_validations.dataValidation = []
    for j,k in enumerate(refcols,1):
        dv = DataValidation(type="list", formula1=f"Reference!${L(j)}$2:${L(j)}${len(rich[k])+1}", allow_blank=True)
        ws.add_data_validation(dv); dv.add(f"{L(col[k])}2:{L(col[k])}1000")
    ws.cell(2, col["location"]).value = rich["location"][0]
    wb.save(sys.argv[sys.argv.index("--rich")+1])
    print("rich ok")
